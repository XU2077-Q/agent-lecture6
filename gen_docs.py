# -*- coding: utf-8 -*-
"""生成两份交付 Word：第六讲-讲解.docx、Agent功能说明.docx（docs/ 与桌面各一份）"""
import shutil
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

DESKTOP = r"C:\Users\Lenovo\Desktop"


def set_cn_fonts(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _cn(run, latin="Calibri"):
    """设中西文字体；先设 latin 字体让 rPr/rFonts 存在，再补 eastAsia"""
    run.font.name = latin
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
        _cn(r)
    return p


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    _cn(run)
    return para


def bullets(doc, items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        _cn(para.add_run(it))


def code(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    para.paragraph_format.left_indent = Pt(18)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Light Grid Accent 1"
    for j, htext in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = htext
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                _cn(run)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = val
            for para in cell.paragraphs:
                for run in para.runs:
                    _cn(run)
    return t


def save(doc, name):
    os.makedirs("docs", exist_ok=True)
    path = os.path.join("docs", name)
    doc.save(path)
    shutil.copy(path, os.path.join(DESKTOP, name))
    print("已生成:", path, "->", DESKTOP)


# ==================================================================
def build_jiangjie():
    doc = Document()
    set_cn_fonts(doc)
    doc.add_heading("第六讲《组装自己的Agent》讲解", 0)

    h(doc, "一、一句话总览")
    p(doc, "前三讲分别做好了三块零件——外部通讯（微信/飞书收发消息）、记忆管理（短期+长期）、"
           "工具调用（tools 协议）；第六讲把它们按八个步骤总装成一个完整的私人 Agent："
           "住在你的电脑上、通过飞书和你聊天、自己有记忆、会调四个工具、按 AGENTS.md 的规矩和 "
           "skills/ 里的技能干活，最后把做好的文件直接发回给你。")

    h(doc, "二、总装架构（P3）")
    bullets(doc, [
        "飞书 App ⇄ 飞书服务器：长连接收消息、REST 回消息、收发文件。",
        "飞书服务器 ⇄ 你的 Python 程序：事件推送进来，程序处理后回发。",
        "Python 程序 ⇄ 大模型：带着工具清单（tools）去问，模型回答或要求调工具。",
        "底座四件套：memory/（短期+长期记忆）、四个工具（run_cmd/read_file/edit_file/send_file）、"
        "AGENTS.md + skills/（工作规矩与技能）、workspace/（收文件、放成品）。",
    ])

    h(doc, "三、八步总装逐页讲解（P4~P19）")
    table(doc,
          ["步骤", "课件页", "做什么", "关键代码点"],
          [
              ["① 接入飞书", "P4~P5",
               "自建应用拿 App ID/Secret，长连接收消息，先做「复读机」打通链路",
               "on_message_receive 里先判 sender_type，忽略机器人自己的消息——防死循环"],
              ["② 接入大模型", "P6~P7",
               "OpenAI 兼容接口三项配置：地址、密钥、型号；ask() 一问一答",
               "POST {base}/chat/completions，Bearer 密钥，取 choices[0].message"],
              ["③ 短期记忆", "P8~P9",
               "每个会话一份 messages，让模型记得本轮聊过什么",
               "滑动窗口只留最近 20 条；落盘 memory/{chat_id}.json，重启可恢复"],
              ["④ 长期记忆", "P10~P11",
               "每天聊过的事追加进日记文件，跨天也能查到",
               "[HH:MM] 角色: 内容 追加 memory/日期.txt；只写不读，读交给工具"],
              ["⑤ 读取记忆", "P12~P13",
               "给模型第一个工具 run_cmd，让它自己 grep 记忆目录",
               "handle() 升级为 Agent 循环：最多连续 10 轮工具调用，防失控"],
              ["⑥ 读写文件", "P14~P15",
               "在 run_cmd 之外给两个文件专用工具",
               "read_file 截 8000 字；edit_file 只替换唯一一段原文，省 token"],
              ["⑦ 发送文件", "P16~P17",
               "把发文件也封装成模型可调用的工具，成品直接发回飞书",
               "先 /im/v1/files 上传拿 file_key 再发 file 消息；CURRENT_CHAT 记当前会话，配合锁一次只干一件事"],
              ["⑧ 装上技能", "P18~P19",
               "AGENTS.md 写工作规矩；skills/ 一技能一文件夹",
               "索引（name+description）进系统提示词，全文任务相关时才 read_file——渐进披露省 token"],
          ])
    p(doc, "P15 页边的思考题「为什么不只给 run_cmd？」：cmd 什么都能干，但文件读写是高频动作，"
           "专用工具有三个好处——参数结构清晰（路径/旧文/新文分开，不用模型拼 shell 引号）、"
           "行为可控可审计（日志一眼看清读了哪个文件）、失败提示可读（edit_file 会明确说"
           "「原文没找到或出现多次」而不是一串 shell 报错）。", bold=True)

    h(doc, "四、三个延伸方向（P20~P22）")
    bullets(doc, [
        "安全（P20）：模型可能幻觉出 rm -rf / 这类危险命令。三道闸——命令黑白名单、"
        "工作区目录限制、危险操作先向用户确认。",
        "定时任务（P21）：现在的 Agent 只能被动回复；把任务登记给调度器，到点触发后走同一条 "
        "Agent 主循环，结果主动推送。Claude Code 的 CronCreate、Kimi 的定时提醒都是这个做法。",
        "更丰富的输入输出（P22）：输入侧接入图片、语音消息；输出侧用飞书消息卡片做流式输出和排版。",
    ])

    h(doc, "五、课后任务与完成情况（P23）")
    p(doc, "作业原文：「让自己的 Agent 做一件事——在飞书或微信里，让自己的 Agent 完成一件具体的事；"
           "例：为喜欢的明星生成一份 PPT，发回给你。截图提交：交了就有分，重在实践，不要抄，不要卷。」")
    p(doc, "完成方式：按课件八步组装出 agent.py（含飞书长连接与 CLI 双模式，同一条 handle() 主循环），"
           "并装了一个 ppt-maker 技能（skills/ppt-maker/SKILL.md）。在 CLI 模式实跑 DeepSeek，"
           "对 Agent 说：「帮我为周杰伦做一份中文PPT，主题是他的音乐生涯，6页左右，保存到 workspace "
           "并把文件发给我」。")
    p(doc, "实跑共两轮，正好把课件两个知识点跑成了真事故与真修复：", bold=True)
    bullets(doc, [
        "第一轮：Agent 读 SKILL.md→写生成脚本→跑脚本→报错→改→再跑……连转 10 轮还没干完，"
        "撞上 P13 的熔断：「工具调用太多次了，我先歇会儿」。查因：技能文件里给的中文字体写法"
        "是 python-docx 的 rFonts 写法，python-pptx 里根本不存在，模型被错误示例带进了死胡同。",
        "修复：把实测可用的字体代码（get_or_add_latin + a:ea 元素）写进 SKILL.md——"
        "这正是 P18「技能全文要写得准」的教训。",
        "第二轮：全流程一次走通——读技能→edit_file 写脚本→run_cmd 执行→send_file 发回→"
        "还主动用 run_cmd 把成品信息追加进了长期记忆。共 10 次工具调用，成品 "
        "workspace/jay_chou_music_career.pptx 六页 16:9，微软雅黑中文字体已嵌入。",
    ])
    p(doc, "留痕：docs/demo_transcript.txt（成功轮全程日志）、"
           "docs/demo_round1_hit_10round_cap.txt（撞熔断的第一轮）。")

    h(doc, "六、术语解释")
    table(doc, ["术语", "解释"], [
        ["长连接（WebSocket）", "客户端与服务器保持一条不断开的通道，服务器有新消息立刻推下来；"
         "飞书机器人用它收消息，不需要自己开公网端口、配 Webhook。"],
        ["Webhook", "反向回调：服务器有事件时 HTTP  POST 到你提供的公网地址。长连接是它的替代方案，"
         "适合没有公网 IP 的本机开发。"],
        ["tenant_access_token", "飞书自建应用的企业级身份凭证，App ID + Secret 换来，"
         "调用发消息、传文件等 REST API 时放在 Authorization 头里，有有效期需缓存复用。"],
        ["OpenAI 兼容接口", "各家大模型厂商照 OpenAI 的 API 形状（/chat/completions、messages 结构）"
         "提供服务，换厂商只改地址、密钥、型号三项，代码不动。"],
        ["messages / role", "与大模型对话的消息列表；role 有 system（规矩）、user（用户）、"
         "assistant（模型）、tool（工具结果）四种。"],
        ["系统提示词（system prompt）", "messages 最前面的一条 system 消息，给模型立规矩、"
         "交代工具与记忆怎么用，本次把 AGENTS.md 和技能索引也装在这里。"],
        ["短期记忆", "当前会话的 messages 历史，让模型记得本轮聊过什么；课件用滑动窗口只留最近 20 条。"],
        ["滑动窗口", "只保留最近 N 条记录的截断策略，防止对话变长后 token 爆炸。"],
        ["长期记忆", "按天追加的日记文件（memory/日期.txt），只写不读；要回想时用 run_cmd 跑 "
         "grep/findstr 去搜。"],
        ["工具调用（tool_calls）", "模型不直接回答，而是在 message 里返回要调用的函数名和参数 JSON；"
         "程序执行后把结果以 role=tool 的消息回填，再问模型。"],
        ["JSON Schema", "描述工具参数结构的声明格式（有哪些参数、类型、是否必填），"
         "登记在 tools 清单里随请求发给模型。"],
        ["Agent 循环", "「问模型→要工具就执行并回填→再问」的循环，直到模型直接回答；"
         "课件设上限 10 轮防失控。"],
        ["复读机", "接入飞书后的第一个最小验证：把用户的话原样回一遍，证明收发链路通了。"],
        ["死循环防护", "机器人收到自己发的消息又触发回调会无限循环；判 sender_type 忽略自己的消息即可。"],
        ["AGENTS.md", "项目根目录的工作规矩文件，启动时读进系统提示词；要短，只写规矩不写技能细节。"],
        ["Skill（技能）", "skills/ 目录下一个文件夹一项技能，里面 SKILL.md 写清什么时候用、怎么做；"
         "用户也可以给 Agent 发新技能。"],
        ["渐进披露", "技能只把 name+description 索引放进系统提示词，全文等任务相关时才 read_file，"
         "省 token——Claude Code 的 Skills 同款设计。"],
        ["workspace", "工作区目录：用户发来的文件先下载到这里，Agent 的成品也放这里，便于管理与限制。"],
        ["file_key", "飞书文件上传成功后返回的标识，发文件消息时凭它引用已上传的文件。"],
        ["CURRENT_CHAT 与锁", "send_file 需要知道发给谁：用全局变量记录正在服务的会话，"
         "配合互斥锁保证一次只处理一个任务。"],
        ["幻觉（hallucination）", "模型编造看似合理实则错误的内容，例如生成根本不存在的命令参数——"
         "安全闸要防的正是它。"],
        ["黑白名单", "对 run_cmd 的命令做过滤：黑名单禁掉 rm -rf 等危险命令，白名单只放行明确安全的。"],
        ["沙箱/工作区限制", "只许 Agent 读写指定目录，目录之外对它「不可见」，把破坏面锁死在沙箱里。"],
        ["人工确认（human-in-the-loop）", "危险操作执行前先发消息问用户，人点了头才动手。"],
        ["调度器（scheduler）", "定时任务的发动机：任务登记在它那里，到点触发，"
         "触发后走同一条 Agent 主循环再主动推送结果。"],
        ["消息卡片", "飞书的富文本消息格式，支持按钮、排版和流式输出，让回答更精致。"],
        ["流式输出", "不等全部生成完，边生成边展示，用户体验上「打字机」式出现。"],
        ["token", "大模型计费与长度单位；省 token 是滑动窗口、8000 字截断、唯一段替换、"
         "渐进披露这些设计的共同动机。"],
    ])

    save(doc, "第六讲-讲解.docx")


# ==================================================================
def build_agent_spec():
    doc = Document()
    set_cn_fonts(doc)
    doc.add_heading("Agent 功能说明（第六讲课后作业）", 0)

    h(doc, "一、这是什么")
    p(doc, "一个按课件八步总装的私人 Agent（agent.py，约 300 行 Python）：住在用户的 Windows 电脑上，"
           "配好飞书凭据时通过飞书长连接聊天，没有凭据时在终端 CLI 模式运行——两种模式走同一条 "
           "handle() Agent 主循环。它有大模型大脑、短期+长期双记忆、四个工具和一套技能系统，"
           "能把做好的文件直接发回给用户。")

    h(doc, "二、架构与运行模式")
    table(doc, ["层", "组成", "说明"], [
        ["通讯层", "飞书长连接 / 终端 CLI", "lark-oapi WebSocket 收消息，REST 回消息发文件；"
         "CLI 模式输入输出在终端，逻辑完全复用"],
        ["大脑", "OpenAI 兼容大模型", "默认 DeepSeek（deepseek-chat），地址/密钥/型号三项可换"],
        ["记忆", "memory/", "短期：{chat_id}.json 滑动窗口 20 条；长期：日期.txt 逐日追加"],
        ["工具", "run_cmd / read_file / edit_file / send_file", "JSON Schema 登记，"
         "tool_calls 协议执行回填"],
        ["规矩与技能", "AGENTS.md + skills/*/SKILL.md", "索引进系统提示词，全文渐进披露"],
        ["工作区", "workspace/", "用户文件与成品产出统一放这里"],
    ])

    h(doc, "三、四个工具")
    table(doc, ["工具", "参数", "行为", "课件出处"], [
        ["run_cmd", "cmd", "执行命令（Windows 走 cmd），输出截 8000 字符；自动把当前 Python "
         "解释器目录前置到 PATH，避开 WindowsApps 的 python 占位 stub", "P12~P13"],
        ["read_file", "path", "读文件内容，最多 8000 字符；文件不存在返回明确提示", "P14~P15"],
        ["edit_file", "path, old_text, new_text", "old_text 留空=新建或整文件重写；否则要求旧文"
         "唯一出现才替换，失败时提示先 read_file 多贴上下文", "P14~P15"],
        ["send_file", "path", "飞书模式：先上传拿 file_key 再发文件消息；CLI 模式：报告成品绝对路径",
         "P16~P17"],
    ])

    h(doc, "四、Agent 主循环")
    code(doc,
         "handle(chat_id, text):\n"
         "  加锁（一次只处理一个任务，CURRENT_CHAT 记当前会话）\n"
         "  载入短期记忆 → 追加用户消息 → 写长期记忆\n"
         "  for 最多 10 轮:\n"
         "      问大模型（带 tools 清单）\n"
         "      若模型直接回答 → 存记忆、写日志、回复，结束\n"
         "      否则逐个执行 tool_calls，结果以 role=tool 回填再问\n"
         "  超 10 轮 → 熔断回复「工具调用太多次了，我先歇会儿」")
    p(doc, "10 轮熔断不是摆设：实跑第一轮就因技能文件里的错误字体示例让模型反复修脚本，"
           "正好触发熔断保护（留痕 docs/demo_round1_hit_10round_cap.txt）。")

    h(doc, "五、记忆系统")
    bullets(doc, [
        "短期记忆：每个 chat_id 一份 messages，滑动窗口只留最近 20 条，落盘 memory/{chat_id}.json，"
        "重启后可恢复。",
        "长期记忆：每天追加 memory/日期.txt，格式 [HH:MM] 角色: 内容，只写不读。",
        "回想方式：模型自己用 run_cmd 跑 findstr/grep 搜 memory/ 目录（系统提示词里写明）。",
    ])

    h(doc, "六、技能系统")
    bullets(doc, [
        "AGENTS.md：六条工作规矩（说中文、成品进 workspace、做完必 send_file、最小改动、"
        "拒危险命令、先查记忆再答），启动时读进系统提示词。",
        "skills/ 一技能一文件夹，SKILL.md 开头 name:/description: 两行做索引进系统提示词；"
        "任务相关时模型才 read_file 读全文——渐进披露省 token。",
        "自带 ppt-maker 技能：用 python-pptx 生成中文 PPT，含实测可用的中文字体代码"
        "（get_or_add_latin + a:ea 写法），脚本与成品都留 workspace/。",
    ])

    h(doc, "七、配置项")
    table(doc, ["配置", "文件 / 环境变量", "说明"], [
        ["大模型地址", "llm_base.txt / LLM_BASE_URL", "默认 https://api.deepseek.com/v1"],
        ["大模型密钥", "llm_key.txt / LLM_API_KEY", "绝不入库（.gitignore）"],
        ["模型型号", "llm_model.txt / LLM_MODEL", "默认 deepseek-chat"],
        ["飞书 App ID", "feishu_app_id.txt / FEISHU_APP_ID", "可选；不配则自动 CLI 模式"],
        ["飞书 Secret", "feishu_app_secret.txt / FEISHU_APP_SECRET", "可选，绝不入库"],
    ])
    p(doc, "读取优先级：环境变量 > 同目录文件 > 默认值。")

    h(doc, "八、安全设计（对应课件 P20）")
    bullets(doc, [
        "AGENTS.md 明文规矩：删除、格式化、关机类命令一律拒绝执行并向用户说明。",
        "工具面收敛：文件读写只走 read_file/edit_file 两个专用工具，行为可审计。",
        "一次一个任务：全局锁 + CURRENT_CHAT，避免并发任务把文件发错会话。",
        "熔断：连续工具调用上限 10 轮，异常时回复报错而非崩溃。",
        "课件延伸的三道闸（命令黑白名单、工作区目录硬限制、危险操作人工确认）留作下一步加强项。",
    ])

    h(doc, "九、实跑演示（课后任务）")
    p(doc, "CLI 模式实跑 DeepSeek，任务：为周杰伦做一份音乐生涯中文 PPT 并发回。"
           "第二轮一次走通：读 SKILL.md → edit_file 写生成脚本 → run_cmd 执行 python-pptx → "
           "send_file 发回 → 主动把成品信息写进长期记忆，共 10 次工具调用。"
           "成品 workspace/jay_chou_music_career.pptx（6 页 16:9，微软雅黑），"
           "全程日志 docs/demo_transcript.txt。")

    h(doc, "十、扩展方向（对应课件 P21~P22）")
    bullets(doc, [
        "定时任务：登记任务给调度器，到点触发同一条 handle() 主循环并主动推送。",
        "图片/语音输入：飞书把多媒体消息推下来，下载后交模型处理。",
        "消息卡片输出：流式 + 排版，替代纯文本回复。",
        "更多技能：往 skills/ 加一个文件夹就是一次升级，无需改主程序。",
    ])

    save(doc, "Agent功能说明.docx")


if __name__ == "__main__":
    build_jiangjie()
    build_agent_spec()
